"""
Document Processor - Handles multiple document formats
"""
import os
import csv
import json
from typing import List, Dict
from PyPDF2 import PdfReader
try:
    from docx import Document
except ImportError:
    Document = None
try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import pandas as pd
except ImportError:
    pd = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None
try:
    import markdown
except ImportError:
    markdown = None


class DocumentProcessor:
    """Processes and extracts text from various document formats"""
    
    def __init__(self):
        self.supported_formats = [
            '.pdf', '.txt', '.docx', '.csv', '.json', 
            '.xlsx', '.xls', '.pptx', '.html', '.htm', '.md'
        ]
    
    def load_document(self, file_path: str) -> str:
        """
        Load and extract text from a document
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Extracted text content
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_ext}. Supported: {self.supported_formats}")
        
        # Route to appropriate extractor
        extractors = {
            '.pdf': self._extract_from_pdf,
            '.txt': self._extract_from_txt,
            '.docx': self._extract_from_docx,
            '.csv': self._extract_from_csv,
            '.json': self._extract_from_json,
            '.xlsx': self._extract_from_excel,
            '.xls': self._extract_from_excel,
            '.pptx': self._extract_from_pptx,
            '.html': self._extract_from_html,
            '.htm': self._extract_from_html,
            '.md': self._extract_from_markdown,
        }
        
        extractor = extractors.get(file_ext)
        if extractor:
            return extractor(file_path)
        else:
            raise ValueError(f"No extractor available for {file_ext}")
    
    def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file"""
        text = []
        try:
            reader = PdfReader(file_path)
            print(f"📄 PDF has {len(reader.pages)} pages")
            
            for i, page in enumerate(reader.pages, 1):
                try:
                    # Try multiple extraction methods
                    page_text = page.extract_text()
                    
                    # Alternative extraction if first method fails
                    if not page_text or len(page_text.strip()) < 10:
                        # Try with different extraction mode
                        page_text = page.extract_text(extraction_mode="layout")
                    
                    if page_text and page_text.strip():
                        text.append(page_text)
                        print(f"  Page {i}: {len(page_text)} chars")
                    else:
                        print(f"  Page {i}: No text extracted (might be image-based)")
                        
                except Exception as e:
                    print(f"  Page {i}: Error - {str(e)}")
                    continue
            
            full_text = '\n'.join(text)
            print(f"✓ Total extracted: {len(full_text)} characters from {len(text)} pages")
            
            if len(full_text.strip()) < 50:
                print("⚠️ WARNING: Very little text extracted. PDF might be image-based or encrypted.")
            
            return full_text
        except Exception as e:
            raise Exception(f"Error reading PDF: {str(e)}")
    
    def _extract_from_txt(self, file_path: str) -> str:
        """Extract text from TXT file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            # Fallback to different encoding
            with open(file_path, 'r', encoding='latin-1') as f:
                return f.read()
    
    def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        if Document is None:
            raise ImportError("python-docx is required for DOCX files. Install: pip install python-docx")
        
        try:
            doc = Document(file_path)
            paragraphs = []
            
            print(f"📄 DOCX has {len(doc.paragraphs)} paragraphs")
            
            # Extract from paragraphs
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    paragraphs.append(para.text)
                    if i < 3:  # Show first 3 for debugging
                        print(f"  Para {i+1}: {para.text[:100]}")
            
            print(f"✓ Extracted {len(paragraphs)} non-empty paragraphs")
            
            # Also extract text from tables
            print(f"📊 DOCX has {len(doc.tables)} tables")
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)
            
            full_text = '\n'.join(paragraphs)
            print(f"✓ Total DOCX text: {len(full_text)} characters")
            
            if len(full_text.strip()) < 50:
                print("⚠️ WARNING: Very little text extracted from DOCX")
            
            return full_text
        except Exception as e:
            print(f"❌ Error reading DOCX: {str(e)}")
            raise Exception(f"Error reading DOCX: {str(e)}")
    
    def _extract_from_csv(self, file_path: str) -> str:
        """Extract text from CSV file"""
        try:
            text_parts = []
            with open(file_path, 'r', encoding='utf-8') as f:
                csv_reader = csv.reader(f)
                headers = next(csv_reader, None)
                if headers:
                    text_parts.append("Columns: " + ", ".join(headers))
                
                for row in csv_reader:
                    text_parts.append(" | ".join(row))
            
            return '\n'.join(text_parts)
        except Exception as e:
            raise Exception(f"Error reading CSV: {str(e)}")
    
    def _extract_from_json(self, file_path: str) -> str:
        """Extract text from JSON file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            # Convert JSON to readable text
            return self._json_to_text(data)
        except Exception as e:
            raise Exception(f"Error reading JSON: {str(e)}")
    
    def _json_to_text(self, obj, prefix="") -> str:
        """Convert JSON object to readable text"""
        lines = []
        
        if isinstance(obj, dict):
            for key, value in obj.items():
                if isinstance(value, (dict, list)):
                    lines.append(f"{prefix}{key}:")
                    lines.append(self._json_to_text(value, prefix + "  "))
                else:
                    lines.append(f"{prefix}{key}: {value}")
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                if isinstance(item, (dict, list)):
                    lines.append(f"{prefix}Item {i+1}:")
                    lines.append(self._json_to_text(item, prefix + "  "))
                else:
                    lines.append(f"{prefix}- {item}")
        else:
            return str(obj)
        
        return '\n'.join(lines)
    
    def _extract_from_excel(self, file_path: str) -> str:
        """Extract text from Excel file"""
        if pd is None:
            raise ImportError("pandas and openpyxl are required for Excel files. Install: pip install pandas openpyxl")
        
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            text_parts = []
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                text_parts.append(f"\n=== Sheet: {sheet_name} ===")
                text_parts.append(df.to_string(index=False))
            
            return '\n'.join(text_parts)
        except Exception as e:
            raise Exception(f"Error reading Excel: {str(e)}")
    
    def _extract_from_pptx(self, file_path: str) -> str:
        """Extract text from PowerPoint file"""
        if Presentation is None:
            raise ImportError("python-pptx is required for PPTX files. Install: pip install python-pptx")
        
        try:
            prs = Presentation(file_path)
            text_parts = []
            
            for i, slide in enumerate(prs.slides, 1):
                text_parts.append(f"\n=== Slide {i} ===")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)
            
            return '\n'.join(text_parts)
        except Exception as e:
            raise Exception(f"Error reading PPTX: {str(e)}")
    
    def _extract_from_html(self, file_path: str) -> str:
        """Extract text from HTML file"""
        if BeautifulSoup is None:
            raise ImportError("beautifulsoup4 is required for HTML files. Install: pip install beautifulsoup4")
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                html_content = f.read()
            
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return text
        except Exception as e:
            raise Exception(f"Error reading HTML: {str(e)}")
    
    def _extract_from_markdown(self, file_path: str) -> str:
        """Extract text from Markdown file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
            
            # If markdown library is available, convert to HTML then extract text
            if markdown is not None and BeautifulSoup is not None:
                html = markdown.markdown(md_content)
                soup = BeautifulSoup(html, 'html.parser')
                return soup.get_text()
            else:
                # Return raw markdown (still readable)
                return md_content
        except Exception as e:
            raise Exception(f"Error reading Markdown: {str(e)}")
    
    def load_multiple_documents(self, file_paths: List[str]) -> dict:
        """
        Load multiple documents
        
        Args:
            file_paths: List of file paths
            
        Returns:
            Dictionary mapping file names to their content
        """
        documents = {}
        for file_path in file_paths:
            try:
                filename = os.path.basename(file_path)
                content = self.load_document(file_path)
                documents[filename] = content
                print(f"✓ Loaded: {filename} ({len(content)} characters)")
            except Exception as e:
                print(f"✗ Failed to load {file_path}: {str(e)}")
        
        return documents


if __name__ == "__main__":
    # Test the document processor
    processor = DocumentProcessor()
    print("Document Processor initialized")
    print(f"Supported formats: {processor.supported_formats}")
